# -*- coding: utf-8 -*-
"""Build two-sheet ablation xlsx from csv_final/ablation_summary.csv.

Per (dataset, metric) rank: first place bold, second place underlined
(ARI / NMI / ACC independently; full precision; ties all bold, no second when all tied).
Usage:
  python build_ablation_xlsx.py --csv-dir csv_final --out out.xlsx [--filter first26|bio11|all31]
"""
import argparse, csv, statistics
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FIRST26 = ["Baron","Mouse_retina","Limb_Muscle","Melanoma_5K","Quake_10x_Spleen","Quake_Smart-seq2_Lung","worm_neuron_cell",
           "breast_cancer_wisconsin_original","dermatology","hepatitis","ISOLET","birds_bones_and_living_habits","cnae9","ecoli",
           "enron","har","harbermans_survival","heart_attack_analysis_prediction_dataset","htru2","micro-mass",
           "rice_dataset_cammeo_and_osmancik","smoker_condition","statlog_image_segmentation","student_grade",
           "turkish_music_emotion","weather"]
BIO11 = ["Bach","Baron","Guo","Limb_Muscle","Macosko","Melanoma_5K","Quake_10x_Spleen","Quake_Smart-seq2_Lung","Tosches","worm_neuron_cell","Young"]
VARIANTS = ["full","base_scmae","anchor_equal_mask","neighbor_equal","no_edge_modulation",
            "fixed_mix_original_aux","mixed_observation_reconstruction"]
METRICS = ["ari","nmi","acc"]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv-dir", required=True); ap.add_argument("--out", required=True)
    ap.add_argument("--filter", default="first26", choices=["first26","bio11","all31"])
    a = ap.parse_args()
    ds_pool = {"first26": FIRST26, "bio11": BIO11, "all31": None}[a.filter]
    with open(a.csv_dir + "/ablation_summary.csv", encoding="utf-8") as f:
        reader = csv.DictReader(f); summ = list(reader); cols = list(reader.fieldnames)
    data = {}
    for r in summ:
        if ds_pool is None or r["dataset"] in ds_pool:
            data.setdefault(r["dataset"], {})[r["variant"]] = r
    datasets = [d for d in (ds_pool or sorted(data.keys())) if d in data]
    rank = {}
    for ds in datasets:
        rank[ds] = {}
        for m in METRICS:
            vals = [(float(data[ds][v][m + "_mean"]), v) for v in VARIANTS if v in data.get(ds, {})]
            tops = sorted(set(x for x, _ in vals), reverse=True)
            rank[ds][m] = (sorted(v for x, v in vals if x == tops[0]),
                           sorted(v for x, v in vals if x == tops[1]) if len(tops) > 1 else [])
    wb = Workbook(); ws = wb.active; ws.title = "ablation_summary"
    ws.append(cols)
    hfill = PatternFill("solid", fgColor="1F4E79")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center")
    bold_font = Font(bold=True); und_font = Font(underline="single")
    for ds in datasets:
        for v in VARIANTS:
            r = data.get(ds, {}).get(v)
            if r is None: continue
            row = [ds, v, r["n_seeds"]]
            for k in ("ari_mean","ari_std","nmi_mean","nmi_std","acc_mean","acc_std"):
                row.append(round(float(r[k]), 4))
            ws.append(row); idx = ws.max_row
            for m in METRICS:
                first, second = rank[ds][m]; c = ws.cell(row=idx, column=cols.index(m + "_mean") + 1)
                if v in first: c.font = bold_font
                elif v in second: c.font = und_font
    ws.freeze_panes = "A2"
    for i, w in enumerate([42, 38, 8, 10, 10, 10, 10, 10, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws2 = wb.create_sheet("macro_summary")
    ws2.append(["variant","n_datasets","ari_mean","ari_std","nmi_mean","nmi_std","acc_mean","acc_std"])
    for v in VARIANTS:
        vals = {m: [] for m in METRICS}
        for ds in datasets:
            r = data.get(ds, {}).get(v)
            if r:
                for m in METRICS: vals[m].append(float(r[m + "_mean"]))
        if not vals["ari"]: continue
        row = [v, len(vals["ari"])]
        for m in METRICS:
            row.append(round(statistics.mean(vals[m]), 4))
            row.append(round(statistics.stdev(vals[m]), 4) if len(vals[m]) > 1 else 0.0)
        ws2.append(row)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws2.freeze_panes = "A2"
    wb.save(a.out)
    print("saved", a.out, "datasets:", len(datasets))

if __name__ == "__main__":
    main()
