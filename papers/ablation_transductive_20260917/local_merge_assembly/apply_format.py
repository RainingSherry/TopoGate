# -*- coding: utf-8 -*-
"""
Apply ranking highlights to the OFFICIAL final deliverable (independent implementation):
for each dataset row x metric group (ARI/NMI/ACC), rank models by mean (parsed from the
cell text) descending; ties broken by column order (leftmost wins).
  1st place -> bold font ; 2nd place -> single underline.
Values are never modified.
"""
import os
import sys
import re
from openpyxl import load_workbook

WORK = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(WORK)
sys.path.insert(0, WORK)
import spec_canonical as S

FINAL = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "合并汇总_3源_56模型×147数据集_均值±std.xlsx")

wb = load_workbook(FINAL)
ws = wb[S.SHEET1_NAME]
rows = list(ws.iter_rows())

# locate metric column groups
groups = {"ARI": [], "NMI": [], "ACC": []}
for c in rows[0][1:]:
    h = str(c.value)
    for met in groups:
        if h.endswith("_" + met):
            groups[met].append(c.column - 1)  # 0-based column index (excluding dataset col 0)
            break

assert all(len(v) == 56 for v in groups.values()), [len(v) for v in groups.values()]

bold_set = set()
under_set = set()
n_bold = n_under = 0
for r in range(1, len(rows)):
    for met, cols in groups.items():
        entries = []
        for ci in cols:
            cell = rows[r][ci]
            v = cell.value
            if v is None or str(v) == "NA":
                continue
            m = re.match(r"^(-?\d+\.\d{6})(?:±(-?\d+\.\d{6}))?$", str(v).strip())
            if not m:
                raise ValueError("unexpected cell text at row %d col %d: %r" % (r, ci, v))
            entries.append((float(m.group(1)), ci, cell))
        entries.sort(key=lambda t: (-t[0], t[1]))
        if len(entries) >= 1:
            cell1 = entries[0][2]
            cell1.font = cell1.font.copy(bold=True)
            bold_set.add((r, entries[0][1]))
            n_bold += 1
        if len(entries) >= 2:
            cell2 = entries[1][2]
            cell2.font = cell2.font.copy(underline="single")
            under_set.add((r, entries[1][1]))
            n_under += 1

wb.save(FINAL)

out = [
    "applied: bold=%d underline=%d  (147 datasets x 3 metrics = %d)" % (n_bold, n_under, 147 * 3),
    "overlap bold&underline cells: %d" % len(bold_set & under_set),
]
with open(os.path.join(WORK, "apply_format_report.txt"), "w", encoding="utf-8") as fh:
    fh.write("\n".join(out))
print("\n".join(out))
