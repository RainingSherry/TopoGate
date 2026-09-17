# -*- coding: utf-8 -*-
"""Finalize: copy the verified merged file to the source folder root, then verify the copy."""
import os
import sys
import shutil
from openpyxl import load_workbook

WORK = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(WORK)
sys.path.insert(0, WORK)
import spec_canonical as S

FINAL_NAME = "合并汇总_3源_56模型×147数据集_均值±std.xlsx"
src = os.path.join(WORK, "agent_3", S.OUTPUT_FILENAME)   # agent3 (identical to agent1/2/reference)
dst = os.path.join(BASE, FINAL_NAME)

shutil.copy2(src, dst)

# verify the copy cell-by-cell against the reference output
def matrix(path, sheet):
    wb = load_workbook(path, read_only=True, data_only=False)
    ws = wb[sheet]
    return [["" if c.value is None else str(c.value) for c in row] for row in ws.iter_rows()]

diff = 0
for sheet in [S.SHEET1_NAME, S.SHEET2_NAME]:
    a = matrix(dst, sheet)
    b = matrix(os.path.join(WORK, "reference", S.OUTPUT_FILENAME), sheet)
    if len(a) != len(b):
        diff += 1
        continue
    for i in range(len(a)):
        if len(a[i]) != len(b[i]):
            diff += 1
            continue
        for j in range(len(a[i])):
            if a[i][j] != b[i][j]:
                diff += 1

lines = [
    "final file: %s" % dst,
    "size: %d" % os.path.getsize(dst),
    "diff cells vs reference: %d" % diff,
]
with open(os.path.join(WORK, "finalize_report.txt"), "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print("\n".join(lines))
